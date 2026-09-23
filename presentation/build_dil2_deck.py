#!/usr/bin/env python3
"""het68 díl 2 — 50min přednáška, tmavý technický styl (PoC / het68)."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = "/workspace/presentation/het68_dil2_50min.pptx"

# Dark technical palette sampled from the project visualization:
# charcoal field, off-white linework, blue LED accent.
BG = RGBColor(0x0E, 0x11, 0x14)
SURFACE = RGBColor(0x17, 0x1C, 0x21)
SURFACE_2 = RGBColor(0x1E, 0x25, 0x2C)
LINE = RGBColor(0x2E, 0x38, 0x44)
TEXT = RGBColor(0xF4, 0xF7, 0xFA)
MUTED = RGBColor(0x9A, 0xA7, 0xB4)
ACCENT = RGBColor(0x3C, 0xB4, 0xF0)
OK = RGBColor(0x5D, 0xCE, 0xA0)
WARN = RGBColor(0xE7, 0xB1, 0x5A)
DIM = RGBColor(0x6E, 0x7C, 0x8A)

SANS = "Inter"
MONO = "JetBrains Mono"

W = 13.333333
H = 7.5


def rgb(shape_fill, color):
    shape_fill.solid()
    shape_fill.fore_color.rgb = color


def style_run(run, name, size, color, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def add_rect(slide, x, y, w, h, fill, line=None, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    rgb(sh.fill, fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if radius:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def add_text(slide, lines, x, y, w, h, size=18, color=TEXT, bold=False,
             font=SANS, align=PP_ALIGN.LEFT, anchor="t", spacing=2):
    if isinstance(lines, str):
        lines = [lines]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    anchor_map = {"t": "t", "ctr": "ctr", "b": "b"}
    tf._txBody.bodyPr.set("anchor", anchor_map[anchor])
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0 if i == 0 else spacing)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        style_run(run, font, size, color, bold)
    return box


def chrome(slide, kicker, page, pages):
    add_rect(slide, 0, 0, W, H, BG)
    add_rect(slide, 0, 0, W, 0.06, ACCENT)
    add_text(slide, "het68", 0.55, 0.22, 1.6, 0.32, 14, ACCENT, True, MONO)
    add_text(slide, kicker, 2.2, 0.22, 8.4, 0.32, 12, MUTED, False, SANS, PP_ALIGN.RIGHT)
    add_rect(slide, 0.55, 7.12, 12.23, 0.01, LINE)
    add_text(slide, "díl 2  ·  50 minut", 0.55, 7.16, 6, 0.28, 11, DIM, False, SANS, anchor="ctr")
    add_text(slide, f"{page}  /  {pages}", 9.5, 7.16, 3.28, 0.28, 11, DIM, False, MONO,
             PP_ALIGN.RIGHT, anchor="ctr")


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new(prs, kicker, page, pages):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, kicker, page, pages)
    return slide


def title_block(slide, title_lines, y=0.62):
    add_text(slide, title_lines, 0.55, y, 12.2, 0.42 * len(title_lines) + 0.15,
             32, TEXT, True, spacing=0)
    return y + 0.46 * len(title_lines) + 0.12


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Jak slyšet dron dříve, než ho uvidíte"
    prs.core_properties.subject = "het68 díl 2 — 50 minut"
    prs.core_properties.author = "het68"
    pages = 21

    # 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, BG)
    add_rect(s, 0, 0, 0.12, H, ACCENT)
    add_text(s, "het68   ·   díl 2   ·   rok poté", 0.7, 1.35, 10, 0.36, 16, ACCENT, True, MONO)
    add_text(s, ["Jak slyšet dron dříve,", "než ho uvidíte"], 0.7, 1.9, 12, 1.8, 48, TEXT, True, spacing=0)
    add_text(s,
             "Šest MEMS mikrofonů. RP2350 počítá směr. Nody se slyší chirpem.",
             0.7, 4.15, 11, 0.4, 20, MUTED)
    add_text(s, "50 minut   ·   z toho 5 minut živě", 0.7, 6.35, 8, 0.3, 14, DIM, False, MONO)
    add_text(s, "github.com/fedurca/het68", 0.7, 6.75, 8, 0.3, 14, ACCENT, False, MONO)
    notes(s, "Přivítat nové i ty, kdo byli loni. Jedna věta: loni karta ještě nehrála, letos ukazuje směr a domluví se se sousedem. Neopakovat celý první díl.")

    # 2 dnes
    s = new(prs, "CO DNES", 2, pages)
    title_block(s, ["Čtyři zastávky, jeden prototyp"])
    beats = [
        ("01", "Karta", "Šest kanálů, které Linux bere jako obyčejný mikrofon."),
        ("02", "Směr", "Azimut a elevace se počítají přímo na Pico 2."),
        ("03", "Chirp", "Sousední uzel, společný čas a hrubá vzdálenost."),
        ("04", "Vy", "Jak se do vývoje zapojit. Cursor, sonda, sériová řádka."),
    ]
    for i, (num, title, body) in enumerate(beats):
        x = 0.55 + i * 3.15
        add_rect(s, x, 2.35, 3.0, 3.55, SURFACE, LINE, 0.08)
        add_rect(s, x, 2.35, 3.0, 0.08, ACCENT)
        add_text(s, num, x + 0.22, 2.6, 2.5, 0.4, 18, ACCENT, True, MONO)
        add_text(s, title, x + 0.22, 3.15, 2.55, 0.5, 26, TEXT, True)
        add_text(s, body, x + 0.22, 3.85, 2.55, 1.6, 16, MUTED)
    notes(s, "0–4 min včetně úvodu. Říct, že demo je pevných 5 minut na konci, ne „když zbyde čas“. Nováčkům stačí tahle mapa.")

    # 3 scoreboard
    s = new(prs, "LONI  →  LETOS", 3, pages)
    title_block(s, ["Co se za rok stalo"])
    rows = [
        ("6kanálová zvukovka", "Linux ji bere. Žádný vlastní ovladač.", "HOTOVO", OK),
        ("Tři nody a triangulace", "Chirp běží. Trojúhelník je otevřená práce.", "OTEVŘENO", WARN),
        ("Telemetrie k anotaci", "Remote ID a časované detekce.", "ROZBĚHNUTO", OK),
        ("Detekce přímo na MCU", "Směr v reálném čase na Pico 2.", "HOTOVO", OK),
        ("Indukční napájení", "Letos ne. Není to tahle přednáška.", "NE", DIM),
    ]
    add_text(s, "LONI", 0.7, 1.85, 4.5, 0.28, 12, DIM, True, MONO)
    add_text(s, "LETOS", 6.15, 1.85, 4.5, 0.28, 12, DIM, True, MONO)
    for i, (left, right, tag, col) in enumerate(rows):
        y = 2.25 + i * 0.9
        add_rect(s, 0.55, y, 5.35, 0.78, SURFACE, LINE, 0.1)
        add_rect(s, 6.05, y, 6.7, 0.78, SURFACE, LINE, 0.1)
        add_text(s, left, 0.75, y + 0.18, 4.95, 0.45, 16, TEXT, True, anchor="ctr")
        add_text(s, right, 6.25, y + 0.18, 4.35, 0.45, 15, TEXT, False, anchor="ctr")
        add_text(s, tag, 10.7, y + 0.22, 1.85, 0.36, 11, col, True, MONO, PP_ALIGN.RIGHT, anchor="ctr")
    notes(s, "Tohle je emocionální háček pro ty, kdo byli loni. Loni slide říkal, že se zvukovka v systému ještě nechová jako zvukovka. Dnes ano. Triangulaci nepřehánět: ranging ano, průsečík tří uzlů ještě ne. Indukci zmínit jednou a jít dál.")

    # 4 teze
    s = new(prs, "TEZE", 4, pages)
    add_text(s, "Jedna věta, kterou si odnést", 0.55, 1.15, 12, 0.35, 14, ACCENT, True, MONO)
    add_text(s, ["Směr se počítá", "na mikrokontroleru."], 0.55, 1.7, 12, 1.9, 44, TEXT, True, spacing=0)
    cards = [
        ("bez Linuxu", "ve výpočtu směru"),
        ("bez Raspberry Pi", "jako druhého počítače"),
        ("bez FPGA", "jako akcelerátoru"),
    ]
    for i, (a, b) in enumerate(cards):
        x = 0.55 + i * 4.15
        add_rect(s, x, 4.35, 3.95, 1.85, SURFACE, LINE, 0.08)
        add_text(s, a, x + 0.25, 4.6, 3.45, 0.55, 22, TEXT, True, anchor="ctr")
        add_text(s, b, x + 0.25, 5.25, 3.45, 0.5, 16, MUTED, anchor="ctr")
    notes(s, "4–8 min spolu s dalším slidem. Zdůraznit: Pico 2 je zároveň mikrofon i kompas. Host je laboratoř, ne mozek lokalizace.")

    # 5 three doors
    s = new(prs, "JEDEN STREAM", 5, pages)
    title_block(s, ["Tytéž vzorky, troje dveře"])
    doors = [
        ("core1", "Směr", "Azimut, elevace a jistota. Senzor funguje i bez počítače."),
        ("USB", "Laboratoř", "48 kHz, 24 bitů, šest kanálů. Audacity, arecord, učení."),
        ("prohlížeč", "Firmware", "UF2 na desku bez toolchainu. První den je z toho zvukovka."),
    ]
    for i, (k, t, b) in enumerate(doors):
        x = 0.55 + i * 4.15
        add_rect(s, x, 2.3, 3.95, 4.15, SURFACE, LINE, 0.08)
        add_text(s, k, x + 0.28, 2.55, 3.4, 0.35, 14, ACCENT, True, MONO)
        add_text(s, t, x + 0.28, 3.1, 3.4, 0.6, 28, TEXT, True)
        add_text(s, b, x + 0.28, 4.0, 3.4, 1.8, 18, MUTED)
    notes(s, "Nováček: první den Audacity. Směr a síť už jsou na desce. Web je cesta firmware na čip, ne obchodní model. SaaS a licence vynechat.")

    # 6 hardware
    s = new(prs, "POLE", 6, pages)
    title_block(s, ["Co leží na stole"])
    items = [
        ("Pico 2", "RP2350, dvě jádra, FPU. Proto produkce na Pico 2."),
        ("Šest ICS-43434", "MEMS, digitální I2S. SEL na zemi, jen levý slot."),
        ("Jedny hodiny", "WS a SCK společné pro všechny. Šest datových linek."),
        ("3,3 V", "Grove shield. Pět voltů mikrofon zničí."),
    ]
    for i, (t, b) in enumerate(items):
        col = i % 2
        row = i // 2
        x = 0.55 + col * 6.35
        y = 2.25 + row * 2.15
        add_rect(s, x, y, 6.1, 1.95, SURFACE, LINE, 0.08)
        add_text(s, t, x + 0.3, y + 0.28, 5.5, 0.5, 24, TEXT, True)
        add_text(s, b, x + 0.3, y + 0.95, 5.5, 0.7, 16, MUTED)
    notes(s, "8–13 min pole. Nepřednášet barvy Grove kabelů. Jednou ukázat desku, jestli je v sále.")

    # 7 cube
    s = new(prs, "GEOMETRIE", 7, pages)
    title_block(s, ["Kostka stojící na špičce"])
    add_text(s, "Mikrofon ve středu každé stěny. Tři kolmé osy skrz střed. Azimut i výška, žádná slepá osa.",
             0.55, 1.9, 7.3, 1.1, 20, MUTED)
    facts = [
        ("+35°", "horní trojice, azimuty po 120°"),
        ("−35°", "spodní trojice, posun o 60°"),
        ("384 mm", "hrana. Loni držák měl 512 mm."),
    ]
    for i, (n, b) in enumerate(facts):
        y = 3.25 + i * 1.15
        add_rect(s, 0.55, y, 7.3, 1.02, SURFACE, LINE, 0.1)
        add_text(s, n, 0.8, y + 0.22, 2.3, 0.58, 22, ACCENT, True, MONO, anchor="ctr")
        add_text(s, b, 3.2, y + 0.28, 4.4, 0.5, 16, TEXT, anchor="ctr")
    add_rect(s, 8.15, 1.9, 4.6, 4.55, SURFACE, LINE, 0.08)
    add_text(s, "OKTAEDR", 8.4, 2.1, 4.1, 0.3, 12, ACCENT, True, MONO)
    add_text(s, ["m1", "m3     +     m2", "", "m4     +     m6", "m5"],
             8.4, 2.8, 4.1, 2.2, 20, TEXT, True, MONO, PP_ALIGN.CENTER, spacing=6)
    add_text(s, "0° je mikrofon 1\nsever, osa +X", 8.4, 5.35, 4.1, 0.8, 14, MUTED, align=PP_ALIGN.CENTER)
    notes(s, "Proč na špičce: tři kolmé základny, dobře podmíněný 3D odhad. 0° ukázat rukou na modelu, ať demo sedí. Změna 512 → 384 je kvůli korelačnímu oknu, ne kvůli módě.")

    # 8 quantum
    s = new(prs, "JEDNO ČÍSLO", 8, pages)
    title_block(s, ["Jeden vzorek je 7,15 mm"])
    add_text(s, "Při 48 kHz a 343 m/s. Posun mikrofonu o tloušťku tužky je celý vzorek chyby směru.",
             0.55, 1.85, 12, 0.7, 20, MUTED)
    big = [
        ("7,15 mm", "dráha na jeden vzorek"),
        ("384 mm", "sweet-spot hrany"),
        ("~1°", "hrubé rozlišení na té hraně"),
    ]
    for i, (n, b) in enumerate(big):
        x = 0.55 + i * 4.15
        add_rect(s, x, 2.85, 3.95, 2.15, SURFACE, LINE, 0.08)
        add_text(s, n, x + 0.25, 3.1, 3.45, 0.8, 32, ACCENT, True, MONO, anchor="ctr")
        add_text(s, b, x + 0.25, 4.05, 3.45, 0.6, 16, TEXT, align=PP_ALIGN.CENTER)
    add_text(s, "Větší kostka je přesnější a zároveň měkčí, větrnější a dřív se potká s aliasingem. Počítání na čipu není ten limit.",
             0.55, 5.3, 12.2, 0.9, 18, MUTED)
    notes(s, "Nepřednášet tabulku 128 až 1024. Jedna věta o trade-offu stačí. Rám: otevřený hliník, molitan, guma. Plné stěny stíní a odrážejí.")

    # 9 pipeline
    s = new(prs, "SNÍMÁNÍ", 9, pages)
    title_block(s, ["Hodiny drží PIO, data tahá DMA"])
    steps = [
        ("1", "Hodiny", "Jedna PIO mašina kreslí WS a SCK pro všech šest."),
        ("2", "Bity", "Šest přijímačů běží ve stejném taktu. Žádné čekání na hranu."),
        ("3", "Paměť", "DMA skládá vzorky. CPU na tom nevisí."),
        ("4", "Dveře", "USB rámec ven. Stejný proud do výpočtu směru."),
    ]
    for i, (n, t, b) in enumerate(steps):
        x = 0.55 + i * 3.15
        add_rect(s, x, 2.3, 3.0, 3.9, SURFACE, LINE, 0.08)
        add_text(s, n, x + 0.22, 2.5, 2.5, 0.45, 20, ACCENT, True, MONO)
        add_text(s, t, x + 0.22, 3.15, 2.55, 0.5, 22, TEXT, True)
        add_text(s, b, x + 0.22, 3.85, 2.55, 1.9, 16, MUTED)
    notes(s, "13–19 min snímaní. Insight: wait na GPIO u SCK prohrával se synchronizátorem pinu, proto RX program hodiny zrcadlí cyklus po cyklu. Říct lidsky: bity se berou naslepo ve správný tik, ne ‚až uvidím hranu‘.")

    # 10 six channels
    s = new(prs, "USB", 10, pages)
    title_block(s, ["Proč zrovna šest kanálů"])
    add_rect(s, 0.55, 2.15, 7.5, 4.3, SURFACE, LINE, 0.08)
    lines = [
        ("Full-Speed", "RP2350 má USB 1.1, 12 Mbit/s. Není to High-Speed."),
        ("864 B / ms", "48 vzorků × 6 kanálů × 3 bajty. Strop izochronního paketu je 1023 B."),
        ("24 bitů", "Host vidí S24_3LE, 48 kHz. Produkt: Pico 6ch Microphone 48k/24."),
    ]
    for i, (t, b) in enumerate(lines):
        y = 2.4 + i * 1.25
        add_text(s, t, 0.85, y, 6.9, 0.38, 20, ACCENT, True, MONO)
        add_text(s, b, 0.85, y + 0.4, 6.9, 0.7, 15, TEXT)
    add_rect(s, 8.25, 2.15, 4.5, 4.3, SURFACE, LINE, 0.08)
    add_text(s, "DŮKAZ", 8.5, 2.4, 4.0, 0.3, 12, ACCENT, True, MONO)
    add_text(s, ["arecord -D hw:X,0", "-c 6 -r 48000", "-f S24_3LE", "capture.wav"],
             8.5, 3.0, 4.0, 1.8, 16, TEXT, False, MONO, spacing=6)
    add_text(s, "Standardní ovladač.\nŽádný dongle.", 8.5, 5.15, 4.0, 0.9, 16, MUTED)
    notes(s, "Šest není designový rozmar, je to strop sběrnice po režii. Jedna jizva 60 s: TinyUSB na RP2350 samo nestačilo, Linux jinak kartu pustil a zase zavřel. Opravy šly zpět upstream. To je věta pro přispěvatele.")

    # 11 cores
    s = new(prs, "DVĚ JÁDRA", 11, pages)
    title_block(s, ["Jedno jádro nesmí zmeškat USB"])
    cols = [
        ("core0", "Zvuk a dům", ["USB a I2S", "sériové příkazy", "barometr a čas", "příjem chirpu"]),
        ("core1", "Směr", ["kruhová fronta", "okno 256 vzorků", "report asi 5× za s", "žádné alokace za běhu"]),
    ]
    for i, (name, sub, bullets) in enumerate(cols):
        x = 0.55 + i * 6.4
        add_rect(s, x, 2.2, 6.15, 4.25, SURFACE, LINE, 0.08)
        add_text(s, name, x + 0.35, 2.45, 5.4, 0.4, 14, ACCENT, True, MONO)
        add_text(s, sub, x + 0.35, 2.95, 5.4, 0.5, 26, TEXT, True)
        add_text(s, bullets, x + 0.35, 3.7, 5.4, 2.3, 18, MUTED, spacing=8)
    notes(s, "19–26 min DOA. Fronta je kruh 2048×6 int16. Flash se zapisuje, jen když USB zrovna nehraje. To je kontrakt, ne detail pro sál — patří do poznámky, když se zeptají.")

    # 12 how doa
    s = new(prs, "SMĚR", 12, pages)
    title_block(s, ["Kdo to slyšel dřív"])
    steps = [
        ("Pásmo", "Dron je spojitý šum zhruba 800 Hz až 6 kHz. Vítr je níž a umí směrový odhad zkazit."),
        ("Zpoždění", "Křížová korelace proti nejsilnějšímu mikrofonu. Zlomky vzorku dopočítá parabola."),
        ("Šipka", "Tři čísla z geometrie kostky. Váhy podle jistoty. Vyjde azimut a elevace."),
        ("Druhý zdroj", "První směr se odečte a hledá se další. Najednou umí dva drony, ne dav."),
    ]
    for i, (t, b) in enumerate(steps):
        y = 2.15 + i * 1.15
        add_rect(s, 0.55, y, 12.2, 1.05, SURFACE, LINE, 0.1)
        add_text(s, f"{i+1:02d}", 0.8, y + 0.28, 0.7, 0.5, 18, ACCENT, True, MONO, anchor="ctr")
        add_text(s, t, 1.7, y + 0.14, 3.2, 0.75, 20, TEXT, True, anchor="ctr")
        add_text(s, b, 5.0, y + 0.18, 7.4, 0.7, 15, MUTED, anchor="ctr")
    notes(s, "Lidsky, ne MUSIC. Rychlost zvuku není kouzelná 343: barometr dává teplotu a tlak, vlhkost se zadá. Bez něj 343. Dron je širokopásmový, proto drží jeden korelační vrchol. Úzký tón na velké kostce by lhal.")

    # 13 outputs
    s = new(prs, "CO Z TOHO LEZE", 13, pages)
    title_block(s, ["Směr je produkt. Třída je odhad."])
    bits = [
        ("az  el", "vodorovně a výška, ve stupních"),
        ("conf", "jak moc si korelace věří"),
        ("dron / vítr", "heuristika v pásmech, ne natrénovaný model"),
        ("~200 ms", "lokalizace, ne sledování každého vzorku"),
    ]
    for i, (t, b) in enumerate(bits):
        col = i % 2
        row = i // 2
        x = 0.55 + col * 6.4
        y = 2.25 + row * 2.15
        add_rect(s, x, y, 6.15, 1.95, SURFACE, LINE, 0.08)
        add_text(s, t, x + 0.3, y + 0.35, 5.55, 0.5, 24, ACCENT, True, MONO)
        add_text(s, b, x + 0.3, y + 1.05, 5.55, 0.55, 16, TEXT)
    notes(s, "Čestná věta, ať vás DSP nechytí: druh ptáka nebo auta je skóre. Směr je to, co obhajujete. Těžší učení patří na počítač, který si stáhne WAV. Časované DET až po TIME SYNC.")

    # 14 why chirp
    s = new(prs, "SÍŤ", 14, pages)
    title_block(s, ["Jeden uzel umí prst, ne metr"])
    add_text(s, "Loňský plán byly tři nody a triangulace. Směr z jedné kostky neříká, jak daleko zdroj je.",
             0.55, 1.9, 12.2, 0.8, 20, MUTED)
    cols = [
        ("Stejné uši", "Piezo na desce krátce tikne. Šestice mikrofonů to slyší. Žádné rádio navíc."),
        ("Až osm sousedů", "Každý má své číslo 0 až 7 a vlastní kód v úvodní značce."),
        ("Dvě čísla navíc", "Společný čas a hrubá vzdálenost. Z tří vzdáleností jednou vznikne bod v prostoru."),
    ]
    for i, (t, b) in enumerate(cols):
        x = 0.55 + i * 4.15
        add_rect(s, x, 3.05, 3.95, 3.35, SURFACE, LINE, 0.08)
        add_text(s, t, x + 0.25, 3.3, 3.45, 0.9, 22, TEXT, True)
        add_text(s, b, x + 0.25, 4.35, 3.45, 1.7, 16, MUTED)
    notes(s, "26–33 min chirp. Nesklouznout do PHY. Říct, že triangulace je důvod, proč chirp existuje, a že průsečík ještě není hotový produkt.")

    # 15 70ms
    s = new(prs, "CHIRP", 15, pages)
    title_block(s, ["Tiknutí, ne sedmivteřinové hvízdání"])
    pair = [
        ("dřív", "~7 s", "dlouhý tón na 4 kHz, nepříjemný a pomalý"),
        ("teď", "~70 ms", "značka a přeskakující tóny mezi 3 a 5,4 kHz"),
    ]
    for i, (era, num, b) in enumerate(pair):
        x = 0.55 + i * 6.4
        add_rect(s, x, 2.2, 6.15, 2.7, SURFACE, LINE, 0.08)
        add_text(s, era, x + 0.3, 2.4, 5.5, 0.3, 13, ACCENT, True, MONO)
        add_text(s, num, x + 0.3, 2.8, 5.5, 0.8, 40, TEXT, True, MONO)
        add_text(s, b, x + 0.3, 3.75, 5.5, 0.8, 16, MUTED)
    add_rect(s, 0.55, 5.1, 12.2, 1.35, SURFACE, LINE, 0.1)
    add_text(s, "1 tik  ≈  9 cm", 0.85, 5.3, 4.5, 0.9, 28, ACCENT, True, MONO, anchor="ctr")
    add_text(s, "Rozlišení vzdálenosti na jeden čip. Na triangulaci začátek, ne geodézie. Pořád to je slyšet.",
             5.5, 5.35, 6.9, 0.9, 16, TEXT, anchor="ctr")
    notes(s, "v1.3 byla asi 7 s DSSS, v1.4 je 279 čipů × 250 µs = 70 ms. Čip 250 µs je asi 8,6 cm při 343 m/s. Říct 9 cm. Nepředstírat ticho ani centimetrovou přesnost. Dosah venku je otevřené měření.")

    # 16 test interface
    s = new(prs, "TESTOVACÍ ROZHRANÍ", 16, pages)
    title_block(s, ["Laboratoř je sériová řádka"])
    add_text(s, "115200 baud. Debug Probe, nebo USB sériový port desky.",
             0.55, 1.85, 12, 0.4, 16, MUTED)
    add_rect(s, 0.55, 2.4, 6.3, 4.05, SURFACE, LINE, 0.08)
    add_text(s, "PŘÍKAZY", 0.85, 2.6, 5.7, 0.3, 12, ACCENT, True, MONO)
    cmds = [
        ("LINK ID 1", "číslo uzlu, a tedy jeho kód"),
        ("LINK BEACON", "pošli jedno tiknutí teď"),
        ("LINK", "soused, vzdálenost, posun hodin"),
        ("TIME INFO", "čas z počítače, z RID, nebo akusticky"),
    ]
    for i, (c, d) in enumerate(cmds):
        y = 3.1 + i * 0.75
        add_text(s, c, 0.85, y, 5.7, 0.32, 18, TEXT, True, MONO)
        add_text(s, d, 0.85, y + 0.32, 5.7, 0.3, 13, MUTED)
    add_rect(s, 7.05, 2.4, 5.7, 4.05, SURFACE, LINE, 0.08)
    add_text(s, "CO MÁ PŘIJÍT", 7.35, 2.6, 5.2, 0.3, 12, ACCENT, True, MONO)
    add_text(s, ["LINK node=1", "NODE id=0", "dist_m=1.4", "offset_us=…", "synced=1"],
             7.35, 3.2, 5.1, 2.2, 18, TEXT, False, MONO, spacing=4)
    add_text(s, "Dva Pico na stole.\nJedno pípnutí, jeden řádek.", 7.35, 5.45, 5.1, 0.75, 15, MUTED)
    notes(s, "Tohle je testovací rozhraní. Žádné zvláštní GUI není potřeba: HELP, LINK, LINK BEACON. Když má web sériovou konzoli, jsou to tytéž příkazy. BEACON se na UART sám nevypisuje, tabulku ukáže až LINK. V demu to říct, ať lidé nečekají spam.")

    # 17 cursor
    s = new(prs, "JAK SE TO STAVÍ", 17, pages)
    title_block(s, ["Cursor navrhuje. Sonda rozhoduje."])
    steps = [
        ("1", "Pravidla", "AGENTS.md. V audio cestě žádné alokace a žádné čekání v přerušení."),
        ("2", "Úprava", "Agent v Cursoru. Krátká změna, kterou jde přečíst."),
        ("3", "Překlad", "./build.sh. Když se to nepřeloží, není to hotové."),
        ("4", "Pravda", "Debug Probe nahraje ELF. UART řekne, jestli to žije."),
    ]
    for i, (n, t, b) in enumerate(steps):
        y = 2.15 + i * 1.15
        add_rect(s, 0.55, y, 12.2, 1.05, SURFACE, LINE, 0.1)
        add_text(s, n, 0.8, y + 0.25, 0.6, 0.55, 22, ACCENT, True, MONO, anchor="ctr")
        add_text(s, t, 1.7, y + 0.25, 2.6, 0.55, 22, TEXT, True, anchor="ctr")
        add_text(s, b, 4.5, y + 0.28, 7.9, 0.55, 16, MUTED, anchor="ctr")
    notes(s, "33–39 min. Neříct ‚AI napsala firmware‘. Říct: model bez pravidel by do rutiny vložil alokaci nebo sleep v callbacku. Sonda je rozhodčí, protože USB zvukovka a ladění se na jedné sběrnici perou — od toho je fixdebugger.sh. Pravda je sériový log, ne diff.")

    # 18 join
    s = new(prs, "PŘIDEJTE SE", 18, pages)
    title_block(s, ["První večer nemusí být USB řadič"])
    tasks = [
        ("Bez kompilátoru", "Stáhnout UF2 z releasu, nahrát, otevřít sérii, napsat HELP."),
        ("Dva uzly", "LINK ID, jedno tiknutí, zapsat dist_m. Práh a dosah jsou otevřené."),
        ("Nahrávka", "arecord šesti kanálů. Skript, model, cokoliv nad WAV."),
        ("Se sondou", "Clone, AGENTS.md, ./build.sh, Debug Probe. Malý pull request."),
    ]
    for i, (t, b) in enumerate(tasks):
        col = i % 2
        row = i // 2
        x = 0.55 + col * 6.4
        y = 2.2 + row * 2.2
        add_rect(s, x, y, 6.15, 2.0, SURFACE, LINE, 0.08)
        add_text(s, t, x + 0.3, y + 0.28, 5.55, 0.45, 20, TEXT, True)
        add_text(s, b, x + 0.3, y + 0.9, 5.55, 0.8, 16, MUTED)
    notes(s, "GPLv3, pět desek v releasech, tag 1.5.1. Pozvat konkrétně: dokumentace pole, práh chirpu, kreslení dvou bodů z LINK, učení na WAV. Ne ‚někdo by mohl‘. ‚Tohle je volné a potřebujeme to.‘")

    # 19 forward
    s = new(prs, "KAM DÁL", 19, pages)
    title_block(s, ["Tři otevřené dveře"])
    doors = [
        ("Trojúhelník", "Tři vzdálenosti a tři směry. Z chirpu udělat bod, ne jen souseda."),
        ("Učení na hostu", "WAV z karty, model na počítači, zpátky na čip jen to, co se vejde do 200 ms."),
        ("Prst pro kameru", "Čas, třída a azimut jako JSON. Kamera se má kam otočit."),
    ]
    for i, (t, b) in enumerate(doors):
        y = 2.2 + i * 1.5
        add_rect(s, 0.55, y, 12.2, 1.35, SURFACE, LINE, 0.1)
        add_text(s, f"0{i+1}", 0.8, y + 0.38, 0.7, 0.55, 20, ACCENT, True, MONO, anchor="ctr")
        add_text(s, t, 1.7, y + 0.18, 4.2, 0.95, 22, TEXT, True, anchor="ctr")
        add_text(s, b, 6.1, y + 0.28, 6.3, 0.85, 16, MUTED, anchor="ctr")
    notes(s, "39–43 min. Remote ID jen jako věta: na deskách s Bluetooth umíme srovnat akustický azimut s GPS dronu. Není to hlavní talk. Anti-drone systém neslibovat. Je to levný prst.")

    # 20 demo
    s = new(prs, "ŽIVĚ", 20, pages)
    title_block(s, ["Pět minut. Stejný scénář i jako video."])
    demo = [
        ("0:00", "Karta", "V systému je Pico 6ch Microphone 48k/24."),
        ("0:40", "Směr", "Zvuk ve třech azimutech. Na sérii az a el."),
        ("2:20", "Chirp", "Druhý uzel. LINK BEACON. Řádek dist_m a offset_us."),
        ("4:20", "HELP", "Na plátně příkazy. Tohle si může zkusit každý."),
    ]
    for i, (tm, t, b) in enumerate(demo):
        y = 2.15 + i * 1.15
        add_rect(s, 0.55, y, 12.2, 1.05, SURFACE, LINE, 0.1)
        add_text(s, tm, 0.8, y + 0.28, 1.3, 0.5, 18, ACCENT, True, MONO, anchor="ctr")
        add_text(s, t, 2.3, y + 0.28, 2.2, 0.5, 20, TEXT, True, anchor="ctr")
        add_text(s, b, 4.7, y + 0.28, 7.7, 0.5, 18, MUTED, anchor="ctr")
    notes(s, "Když USB umře, stačí UART a směr. Když umře i UART, pusť video bez omluvy delší než věta. Chirp je křehčí: měj 90s záznam tiknutí. Živě neflashovat. 0° je mikrofon 1.")

    # 21 close
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, BG)
    add_rect(s, 0, 0, 0.12, H, ACCENT)
    add_text(s, "het68", 0.7, 1.2, 6, 0.35, 16, ACCENT, True, MONO)
    add_text(s, ["Zvukovka. Kompas.", "Uzel, který se slyší", "se sousedem."], 0.7, 1.75, 12, 2.4, 40, TEXT, True, spacing=0)
    add_text(s, "První krok je HELP na 115200.", 0.7, 4.7, 10, 0.4, 20, MUTED)
    add_text(s, "github.com/fedurca/het68", 0.7, 5.5, 10, 0.4, 22, ACCENT, True, MONO)
    add_text(s, "release 1.5.1   ·   GPLv3   ·   Debug Probe, až budete překládat",
             0.7, 6.15, 11, 0.35, 14, DIM)
    notes(s, "48–50 min. Zastavit. Nechat odkaz na plátně. Otázky až po čase, pokud je konference tak staví. Poděkovat lidem, kteří chtějí měřit dosah chirpu.")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
